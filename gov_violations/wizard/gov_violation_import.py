import base64
import datetime
import io
import re

from odoo import _, fields, models
from odoo.exceptions import UserError

try:
    import openpyxl
except ImportError:  # pragma: no cover - declared in external_dependencies
    openpyxl = None

# The 20 columns of the Balady master, in order. The template download and
# the import both use this layout; new violations are appended as new rows.
BALADY_COLUMNS = [
    'رقم المخالفة',
    'تاريخ رصد المخالفة',
    'رقم الترخيص البلدي / رقم الهوية',
    'نوع الترخيص البلدي',
    'الأمانة / البلدية',
    'اسم المنشأة',
    'بند اللائحة',
    'عدد مرات تكرار المخالفة',
    'عدد وحدات المخالفة',
    'آخر موعد لتقديم الاعتراض',
    'الأيام المتبقية للاعتراض',
    'قيمة الغرامة (ر.س)',
    'رقم الفاتورة',
    'حالة فاتورة المخالفة',
    'بند المخالفة الرئيسي',
    'ملاحظات المراقب',
    'رقم طلب الاعتراض',
    'تاريخ طلب الاعتراض',
    'حالة طلب الاعتراض',
    'قرار لجنة الاعتراض',
]
DATE_COLUMNS = {1, 9, 17}
NUMBER_COLUMNS = {7, 8, 10, 11}
SCIENTIFIC_RE = re.compile(r'^\d(\.\d+)?[eE][+-]?\d+$')


def cell_text(value):
    """Excel cell -> text, without ever going through a float."""
    if value is None:
        return None
    if isinstance(value, bool):
        return str(value)
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float):
        if value.is_integer():
            return str(int(value))
        return repr(value)
    if isinstance(value, (datetime.datetime, datetime.date)):
        return value.strftime('%Y-%m-%d')
    text = str(value).strip()
    return text or None


def cell_date(value):
    if value is None or value == '':
        return None
    if isinstance(value, datetime.datetime):
        return value.date()
    if isinstance(value, datetime.date):
        return value
    text = str(value).strip()
    for fmt in ('%Y-%m-%d', '%d/%m/%Y', '%Y/%m/%d', '%d-%m-%Y'):
        try:
            return datetime.datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    return None


def cell_number(value):
    if value is None or value == '':
        return 0
    if isinstance(value, (int, float)):
        return value
    try:
        return float(str(value).replace(',', '').strip())
    except ValueError:
        return 0


def read_balady_rows(content):
    """Rows of a Balady master workbook as cleaned 20-value lists.
    The header row is located by its first cell, so title rows above it
    (totals, instructions) are skipped."""
    if openpyxl is None:
        raise UserError(_('مكتبة openpyxl غير مثبتة على الخادم.'))
    workbook = openpyxl.load_workbook(io.BytesIO(content), data_only=True, read_only=True)
    sheet = workbook.worksheets[0]
    rows, warnings, header_found = [], [], False
    for raw in sheet.iter_rows(values_only=True):
        raw = list(raw) + [None] * (20 - len(raw))
        first = cell_text(raw[0])
        if not header_found:
            if first == BALADY_COLUMNS[0]:
                header_found = True
            continue
        if not first:
            continue
        row = []
        for index, value in enumerate(raw[:20]):
            if index in DATE_COLUMNS:
                row.append(cell_date(value))
            elif index in NUMBER_COLUMNS:
                row.append(cell_number(value))
            else:
                text = cell_text(value)
                if text and index in (2, 12) and SCIENTIFIC_RE.match(text):
                    warnings.append(_('%(number)s: رقم بصيغة علمية (%(value)s) - احفظ العمود كنص في Excel',
                                      number=first, value=text))
                row.append(text)
        rows.append(row)
    if not header_found:
        raise UserError(_('لم يتم العثور على صف العناوين (رقم المخالفة) في الملف.'))
    return rows, warnings


class GovViolationImport(models.TransientModel):
    _name = 'gov.violation.import'
    _description = 'استيراد ملف بلدي'

    file = fields.Binary(string='ملف بلدي (xlsx)')
    filename = fields.Char(string='اسم الملف')
    full_register = fields.Boolean(
        string='الملف يمثل السجل الكامل لغير المسدد',
        default=True,
        help='عند التفعيل: المخالفات البلدية غير المسددة في النظام وغير الموجودة في '
             'الملف تُعلَّم "يحتاج تأكيد" وتُستبعد من المبلغ المستحق (لا تُحذف ولا '
             'تُسجَّل كمسددة).',
    )
    state = fields.Selection([('draft', 'draft'), ('done', 'done')], default='draft')
    result = fields.Text(string='النتيجة', readonly=True)

    def action_import(self):
        self.ensure_one()
        if not self.file:
            raise UserError(_('اختر ملف بلدي أولاً.'))
        rows, warnings = read_balady_rows(base64.b64decode(self.file))
        summary = self.env['gov.violation']._gov_import_balady(
            rows, full_register=self.full_register)
        lines = [
            _('صفوف الملف: %s', len(rows)),
            _('مخالفات جديدة: %s', summary['created']),
            _('مخالفات محدثة: %s', summary['updated']),
        ]
        if summary['new_licenses']:
            lines.append(_('تراخيص جديدة بدون فرع: %s', ', '.join(summary['new_licenses'])))
        if summary['missing_from_register']:
            lines.append(_('غير موجودة في الملف (تحتاج تأكيد): %s',
                           ', '.join(summary['missing_from_register'])))
        for error in summary['errors']:
            lines.append(_('خطأ: %(row)s - %(error)s', row=error['row'], error=error['error']))
        lines.extend(warnings)
        self.write({'state': 'done', 'result': '\n'.join(lines)})
        return {
            'type': 'ir.actions.act_window',
            'res_model': self._name,
            'res_id': self.id,
            'view_mode': 'form',
            'target': 'new',
        }

    def action_download_template(self):
        if openpyxl is None:
            raise UserError(_('مكتبة openpyxl غير مثبتة على الخادم.'))
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        sheet.title = 'المخالفات'
        sheet.sheet_view.rightToLeft = True
        sheet.append(BALADY_COLUMNS)
        # Text format for the number-like columns, so Excel keeps every digit.
        for col in (1, 3, 13, 17):
            for row in range(2, 1001):
                sheet.cell(row=row, column=col).number_format = '@'
        buffer = io.BytesIO()
        workbook.save(buffer)
        attachment = self.env['ir.attachment'].create({
            'name': 'قالب_مخالفات_بلدي.xlsx',
            'datas': base64.b64encode(buffer.getvalue()),
            'mimetype': 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        })
        return {
            'type': 'ir.actions.act_url',
            'url': '/web/content/%s?download=true' % attachment.id,
            'target': 'self',
        }

    def action_open_violations(self):
        return self.env['ir.actions.act_window']._for_xml_id('gov_violations.action_gov_violation')
